"""Turn a file of any supported type into plain text for the extraction agent.

This is parsing, not judgment, so it is all code and no LLM. One reader per file type,
chosen by extension, with a registry so an unsupported type fails with a sentence an
operator can act on rather than a UnicodeDecodeError or silent mojibake.

Readers for formats beyond the core text ones import their library lazily. A client who
sends .xlsx invoices installs openpyxl; nobody else pays for the dependency.
"""

from pathlib import Path


class DocumentError(Exception):
    """The file cannot be turned into text. The message is written for a human."""


# Extensions that are already text on disk. Reading them is a decode, not a parse, and the
# extraction agent is perfectly able to read raw JSON, CSV or XML.
PLAIN_TEXT = {".txt", ".json", ".csv", ".xml", ".md", ".tsv", ".yaml", ".yml"}

# Formats that carry invoice data but need a real decoder before the agent sees anything.
NEEDS_OCR = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".heic", ".gif"}


def _looks_like_text(text: str, threshold: float = 0.85) -> bool:
    """True when the decoded string is plausibly a human-readable document."""
    if not text:
        return False
    sane = sum(1 for ch in text if ch.isascii() and (ch.isprintable() or ch in "\n\r\t"))
    return sane / len(text) >= threshold


def _read_plain(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # Windows-authored invoices are commonly cp1252. Try it before giving up, because
        # the alternative is failing on a smart quote.
        try:
            text = path.read_text(encoding="cp1252")
        except UnicodeDecodeError as exc:
            raise DocumentError(
                f"{path.name} is not valid UTF-8 or Windows-1252 text; it may be binary "
                f"with a misleading extension"
            ) from exc
        # cp1252 maps almost every byte, so it decodes binary without ever raising. The
        # failure then shows up as mojibake inside the extraction prompt instead of here.
        # Check the result actually looks like prose before passing it on.
        if not _looks_like_text(text):
            raise DocumentError(
                f"{path.name} decoded to mostly non-text characters; it is probably "
                f"binary with a misleading extension"
            )
        return text


def _read_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise DocumentError("reading PDFs needs pdfplumber: uv add pdfplumber") from exc

    try:
        with pdfplumber.open(path) as pdf:
            pages = [(p.extract_text() or "").strip() for p in pdf.pages]
            has_images = any(p.images for p in pdf.pages)
    except Exception as exc:
        raise DocumentError(f"{path.name} could not be opened as a PDF: {exc}") from exc

    text = "\n\n".join(p for p in pages if p)
    if not text:
        # A PDF with no text layer is a photograph in a wrapper. Say so, because "empty
        # invoice" and "scanned invoice" need completely different responses.
        hint = "it appears to be a scan" if has_images else "it contains no text objects"
        raise DocumentError(
            f"{path.name} has no extractable text layer ({hint}); OCR would be required"
        )
    return text


def _read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentError("reading .xlsx needs openpyxl: uv add openpyxl") from exc

    wb = load_workbook(path, data_only=True)  # data_only: formula results, not formulas
    out = []
    for ws in wb.worksheets:
        out.append(f"--- sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                out.append("\t".join(cells).rstrip())
    return "\n".join(out)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentError("reading .docx needs python-docx: uv add python-docx") from exc

    doc = Document(path)
    out = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # invoice line items usually live in a table, not a paragraph
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append("\t".join(cells))
    return "\n".join(out)


def _read_html(path: Path) -> str:
    import re

    raw = _read_plain(path)
    raw = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)</(tr|p|div|h[1-6]|li)>", "\n", raw)
    raw = re.sub(r"(?i)</t[dh]>", "\t", raw)
    text = re.sub(r"<[^>]+>", "", raw)
    for a, b in (("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'), ("&nbsp;", " ")):
        text = text.replace(a, b)
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


READERS = {
    ".pdf": _read_pdf,
    ".xlsx": _read_xlsx,
    ".xlsm": _read_xlsx,
    ".docx": _read_docx,
    ".html": _read_html,
    ".htm": _read_html,
    **{ext: _read_plain for ext in PLAIN_TEXT},
}


def supported() -> list[str]:
    return sorted(READERS)


def load(path: str | Path) -> tuple[str, str]:
    """Read a document to text.

    Returns (text, reader_name) so the run log records how the bytes were decoded. An
    extraction that goes wrong on a PDF is a different investigation from one that goes
    wrong on a CSV, and the log should not make anyone guess which happened.
    """
    path = Path(path)
    if not path.exists():
        raise DocumentError(f"{path} does not exist")
    if path.stat().st_size == 0:
        raise DocumentError(f"{path.name} is empty")

    ext = path.suffix.lower()
    if ext in NEEDS_OCR:
        raise DocumentError(
            f"{path.name} is an image ({ext}); OCR is not configured, so it needs to be "
            f"handled by a person or converted to a text-bearing PDF"
        )

    reader = READERS.get(ext)
    if reader is None:
        raise DocumentError(
            f"no reader for '{ext or path.name}'; supported types are "
            f"{', '.join(supported())}"
        )

    text = reader(path).strip()
    if not text:
        raise DocumentError(f"{path.name} produced no text")
    return text, reader.__name__.removeprefix("_read_")
